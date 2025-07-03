import re

from DrissionPage import ChromiumPage, SessionPage
from DrissionPage.common import Settings
from DrissionPage.errors import ElementNotFoundError

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.shared import Pt
from urllib.parse import urlparse, parse_qs

kemu = ['1160566','1161128','1160280','1160279']

page = ChromiumPage()
page.get('https://s.zaixiankaoshi.com/student/109895')
print(page.title)
ele = page.ele('@placeholder=请输入您的学员账号')
ele.input("15615198376")
# 定位到密码文本框并输入密码
page.ele('@placeholder=请输入您的学员密码').input("123456")
# 点击登录按钮
page.ele('tag:button').click()

page.wait.load_start()


# document = Document()
# document.styles['Normal'].font.name = u'宋体'
# document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
# document.styles['Normal'].font.size = Pt(11)

for i in range(1, 3):

    document = Document()
    document.styles['Normal'].font.name = u'宋体'
    document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    document.styles['Normal'].font.size = Pt(11)

    url = f'https://s.zaixiankaoshi.com/sctk/'
    page.get(url)
    page.wait.eles_loaded('xpath://*[@id="__layout"]/section/section/main/div/div/div[2]/div[2]/div/div/div[1]/div[3]/table/tbody/tr[1]/td[8]/div/button/span')

    a_title =page.ele('xpath://*[@id="__layout"]/section/section/main/div/div/div[2]/div[2]/div/div/div[1]/div[3]/table/tbody/tr['+str(i)+']/td[2]/div/span').text

    buttons = page.ele('xpath://*[@id="__layout"]/section/section/main/div/div/div[2]/div[2]/div/div/div[1]/div[3]/table/tbody/tr['+str(i)+']/td[8]/div/button')
    buttons.click()

    page.ele('xpath://*[@id="body"]/div/div/div[1]/div[2]/div[3]/div/a[2]').click()
    page.wait.eles_loaded('xpath://*[@id="body"]/div/div[1]/div[2]/div[1]/div[2]/div[1]/div/div[1]/b/text()')

    beit = page.ele('xpath://*[@id="body"]/div/div[1]/div[2]/div[2]/div[3]/p[2]/span[2]/div')
    if 'is-checked' not in beit.attrs['class']:
        page.ele('xpath://*[@id="body"]/div/div[1]/div[2]/div[2]/div[3]/p[2]/span[2]/div/span').click()

    num1 = page.ele('xpath://*[@id="body"]/div/div[1]/div[2]/div[1]/div[1]/div/div[1]/div/span[2]').text
    num2 = num1.split('/')
    num10 = int(re.findall(r'\d+', num2[0])[0])
    num11 = int(re.findall(r'\d+', num2[1])[0])

    while num10 != num11:
        num1 = page.ele('xpath://*[@id="body"]/div/div[1]/div[2]/div[1]/div[1]/div/div[1]/div/span[2]').text
        num2 = num1.split('/')
        num10 = int(re.findall(r'\d+', num2[0])[0])
        num11 = int(re.findall(r'\d+', num2[1])[0])
        title = page.ele('@class=qusetion-box').text
        document.add_paragraph(title)
        option2 = document.add_paragraph('')
        type = page.ele('@class=topic-type').text
        options = page.s_eles('@class^option')
        answer = ''
        for opt in options:
            option2.add_run("\n" + opt.raw_text).bold = True
            if 'right' in opt.attrs['class']:
                answer = answer + opt.raw_text.split(" ")[0] + ','
        option2.add_run("\n" + "正确答案: " + answer[:-1]).bold = True
        page.ele('@@class:el-button el-button--primary el-button--small@@text():下一题', timeout=5).click()
        page.wait(float(1))
    document.save(a_title+'2.docx')