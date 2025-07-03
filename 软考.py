import os
import requests
from DrissionPage import ChromiumPage, SessionPage

from docx import Document
from docx.shared import Inches
from urllib.parse import urlparse, parse_qs, urlencode


def download_image(url):
    name = os.path.basename(url)

    response = requests.get(url)
    if response.status_code == 200:
        with open(name, 'wb') as f:
            f.write(response.content)
        return name
    else:
        raise Exception("无法下载图片，状态码: {}".format(response.status_code))


s1=["高级信息系统项目管理师:55971678f3a6404ba6179b89692829f2", "高级系统分析师:057641fa9f9540b1877ab054adc12b65", "高级系统架构设计师:5a54594b40e24123843cb0cda46eaa8a", "高级网络规划设计师:078908c44df04b3799519b1adf4f1e1d", "高级系统规划与管理师:3c2d17d0a657404f9c57b102bd5635cd", "中级系统集成项目管理工程师:11743810783b404e93394ac3513e1a89", "中级软件设计师:700c85921726418bbeadbe6c2db2763c", "中级网络工程师:6c237c1690ba49f89e978885b6da7b17", "中级信息系统监理师:6090d81481504e55890fd24bc69d2a4c", "中级数据库系统工程师:9ef64188059143ffba9e72c137c67659", "中级软件评测师:19e59669d11c4cf0b15aadcefb4d967a", "中级嵌入式系统设计师:e1f5bda7ebc94c968bc970175a97ccd7", "中级电子商务设计师:1714e4b3bce248859dc27db1281d24d0", "中级多媒体应用设计师:abcceba6ca424b22afc7cf73c6bee3b1", "中级信息系统管理工程师:3b99eaa897d74479af3a1480c6747998", "中级信息安全工程师:57031ec59cd64b9abd5737f2faede4f2", "初级程序员:8487c04e6b7749419a646ffb3d406509", "初级信息处理技术员:e117b6b0cd7f4be9987f165fc7d2eda2", "初级网络管理员:78056250e8624a3e9c3a1d6a459df06f", "初级信息系统运行管理员:3546996e02294ee69261b7f10015cb8d"]

BaseUrl = "https://www.lightsoft.tech/"

page = ChromiumPage()
page.get('https://www.lightsoft.tech/doquestion/subject?doType=0&subjectId=5a54594b40e24123843cb0cda46eaa8a')
b1 = page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[1]/div')
b1.click()
b1.click()

tb1 = page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[2]/div/div[2]/div/div/div[1]/table/tbody').children()
tbs = []
titles = []
for tr1 in tb1:
    ela = tr1.ele("xpath://a")
    tbs.append(ela.attr("href"))
    titles.append(tr1.ele("xpath://th").text)

for indexa, tab in enumerate(tbs):
    flag = True
    numb = 0
    parsed_url = urlparse(tab)
    doc = Document()
    # 使用 parse_qs() 提取查询参数
    query_params = parse_qs(parsed_url.query)
    while flag:
        query_params['index'] = numb
        # 将查询参数重新编码为 URL 编码格式
        encoded_query = urlencode(query_params, doseq=True)
        # 重新构建 URL
        try:
            new_tab = parsed_url.scheme + '://' + parsed_url.netloc + parsed_url.path + '?' + encoded_query
            page.get(new_tab)
            page.wait.eles_loaded('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[1]/div/div',2)
            dchilden = page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[1]/div/div/div').children()
            pis = page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[1]/div/div/div').ele("xpath://p")
            if dchilden and pis:
                for index,timu in enumerate(dchilden):
                     if index == 0:
                        doc.add_paragraph(str(numb + 1) + '、' + timu.text)
                        if "img" in timu.html:
                            img = timu.ele('xpath://img')
                            doc.add_picture(download_image(img.link), width=Inches(4))
                     else:
                        if "img" in timu.html:
                            img = timu.ele('xpath://img')
                            if img:
                                doc.add_picture(download_image(img.link), width=Inches(4))
                            else:
                                doc.add_picture(download_image(timu.link), width=Inches(4))
                        else:
                            doc.add_paragraph(timu.text)
            else:
                doc.add_paragraph(page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[1]/div/div/div').text)

            option2 = doc.add_paragraph('')
            if "综合" in titles[indexa]:
                for xua in page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[2]/div/div/div/div').children():
                    option2.add_run(xua.text + '\n')
                page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[2]/div/div/div/div/div[1]/button').click()
                page.wait.eles_loaded('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[2]/div/div/div[2]')
                dan = page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[2]/div/div/div[2]/div[2]/div[1]')
                option2.add_run(dan.text + '\n')
            elif "案例" in titles[indexa]:
                for xua in page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[2]/div/div/div[1]').children():
                    option2.add_run(xua.text + '\n')
            else:
                option2.add_run(page.ele('xpath://*[@id="__next"]/div[1]/div/div/div[1]/div[2]/div/div/div').text)
            numb = numb +1
        except  Exception as e:
            print(e)
            if "Internal Server Error" in page.html:
                flag = False

    doc.save(titles[indexa]+".doc")

